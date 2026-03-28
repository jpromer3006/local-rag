#!/usr/bin/env python3
"""Core utilities for local RAG over centralized files using SQLite FTS5 + Ollama."""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sqlite3
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

# Optional parsers
try:
    import pdfplumber  # type: ignore
except Exception:
    pdfplumber = None

try:
    import docx  # type: ignore
except Exception:
    docx = None

try:
    from pptx import Presentation  # type: ignore
except Exception:
    Presentation = None


SUPPORTED_TEXT_EXTS = {
    ".txt", ".md", ".csv", ".json", ".jsonl", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".sql", ".log",
    ".yaml", ".yml", ".toml", ".ini", ".sh", ".zsh",
}

SKIP_DIR_NAMES = {
    ".git",
    "node_modules",
    "__pycache__",
    "venv",
    ".venv",
    "dist",
    "build",
}


@dataclass
class Chunk:
    path: str
    rel_path: str
    chunk_index: int
    content: str
    file_sha256: str
    category: str
    bytes_size: int
    modified_ts: float


def load_config(config_path: Path) -> Dict[str, object]:
    return json.loads(config_path.read_text(encoding="utf-8"))


def category_for_path(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".pdf", ".doc", ".docx", ".ppt", ".pptx", ".txt", ".md", ".csv"}:
        return "documents"
    if ext in {".py", ".js", ".ts", ".java", ".go", ".rs", ".sql", ".sh", ".zsh"}:
        return "code"
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"}:
        return "images"
    return "other"


def is_excluded(path: Path, excluded_roots: Sequence[Path]) -> bool:
    abs_path = path.resolve()
    for root in excluded_roots:
        root_res = root.resolve()
        if abs_path == root_res or str(abs_path).startswith(str(root_res) + os.sep):
            return True
    return False


def iter_candidate_files(
    source_paths: Sequence[Path],
    exclude_paths: Sequence[Path],
    allowed_extensions: Sequence[str],
    max_files: int,
) -> Iterable[Path]:
    allowed = {x.lower() for x in allowed_extensions}
    emitted = 0

    for root in source_paths:
        if not root.exists():
            continue
        if root.is_file():
            if root.suffix.lower() in allowed and not is_excluded(root, exclude_paths):
                yield root.resolve()
            continue

        for current_root, dirs, files in os.walk(root, topdown=True):
            current_path = Path(current_root).resolve()
            dirs[:] = [
                d for d in dirs
                if (
                    not is_excluded(current_path / d, exclude_paths)
                    and d not in SKIP_DIR_NAMES
                    and not d.endswith(".app")
                )
            ]
            for name in files:
                candidate = (current_path / name).resolve()
                if is_excluded(candidate, exclude_paths):
                    continue
                if candidate.suffix.lower() not in allowed:
                    continue
                yield candidate
                emitted += 1
                if max_files > 0 and emitted >= max_files:
                    return


def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_pdf(path: Path) -> str:
    if pdfplumber is None:
        return ""
    parts: List[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    return "\n\n".join(parts)


def read_docx(path: Path) -> str:
    if docx is None:
        return ""
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def read_pptx(path: Path) -> str:
    if Presentation is None:
        return ""
    prs = Presentation(str(path))
    slides: List[str] = []
    for i, slide in enumerate(prs.slides):
        lines: List[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        lines.append(text)
        if lines:
            slides.append(f"[Slide {i + 1}]\n" + "\n".join(lines))
    return "\n\n".join(slides)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in SUPPORTED_TEXT_EXTS:
        return read_text_file(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    if ext == ".pptx":
        return read_pptx(path)
    return ""


def file_sha256(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            chunk = f.read(chunk_size)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def chunk_text(text: str, chunk_size: int, overlap: int) -> List[str]:
    clean = " ".join(text.split())
    if not clean:
        return []
    if overlap >= chunk_size:
        overlap = max(0, chunk_size // 4)
    step = max(1, chunk_size - overlap)
    chunks: List[str] = []
    for start in range(0, len(clean), step):
        chunk = clean[start:start + chunk_size]
        if chunk.strip():
            chunks.append(chunk)
        if start + chunk_size >= len(clean):
            break
    return chunks


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA journal_mode=WAL;
        PRAGMA synchronous=NORMAL;

        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            path TEXT NOT NULL,
            rel_path TEXT NOT NULL,
            chunk_index INTEGER NOT NULL,
            content TEXT NOT NULL,
            file_sha256 TEXT NOT NULL,
            category TEXT NOT NULL,
            bytes_size INTEGER NOT NULL,
            modified_ts REAL NOT NULL,
            ingested_at TEXT NOT NULL
        );

        CREATE UNIQUE INDEX IF NOT EXISTS idx_chunks_unique
        ON chunks(path, chunk_index, file_sha256);

        CREATE INDEX IF NOT EXISTS idx_chunks_path ON chunks(path);
        CREATE INDEX IF NOT EXISTS idx_chunks_sha ON chunks(file_sha256);

        CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
            content,
            rel_path,
            path,
            category,
            content='chunks',
            content_rowid='id'
        );

        CREATE TRIGGER IF NOT EXISTS chunks_ai AFTER INSERT ON chunks BEGIN
            INSERT INTO chunks_fts(rowid, content, rel_path, path, category)
            VALUES (new.id, new.content, new.rel_path, new.path, new.category);
        END;

        CREATE TRIGGER IF NOT EXISTS chunks_ad AFTER DELETE ON chunks BEGIN
            INSERT INTO chunks_fts(chunks_fts, rowid, content, rel_path, path, category)
            VALUES('delete', old.id, old.content, old.rel_path, old.path, old.category);
        END;

        CREATE TRIGGER IF NOT EXISTS chunks_au AFTER UPDATE ON chunks BEGIN
            INSERT INTO chunks_fts(chunks_fts, rowid, content, rel_path, path, category)
            VALUES('delete', old.id, old.content, old.rel_path, old.path, old.category);
            INSERT INTO chunks_fts(rowid, content, rel_path, path, category)
            VALUES (new.id, new.content, new.rel_path, new.path, new.category);
        END;
        """
    )
    conn.commit()


def upsert_file_chunks(
    conn: sqlite3.Connection,
    path: Path,
    rel_path: str,
    chunks: Sequence[str],
    digest: str,
    category: str,
    size: int,
    mtime: float,
) -> Tuple[int, int]:
    # Replace existing chunks when file hash changed or path already indexed.
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM chunks WHERE path = ?", (str(path),))
    existing = int(cur.fetchone()[0])

    if existing:
        cur.execute("DELETE FROM chunks WHERE path = ?", (str(path),))

    inserted = 0
    now = dt.datetime.utcnow().isoformat(timespec="seconds") + "Z"
    for i, chunk in enumerate(chunks):
        cur.execute(
            """
            INSERT OR IGNORE INTO chunks
            (path, rel_path, chunk_index, content, file_sha256, category, bytes_size, modified_ts, ingested_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (str(path), rel_path, i, chunk, digest, category, size, mtime, now),
        )
        inserted += cur.rowcount

    conn.commit()
    return existing, inserted


def search_chunks(
    conn: sqlite3.Connection,
    query: str,
    top_k: int,
) -> List[Dict[str, object]]:
    terms = [t for t in re.findall(r"[A-Za-z0-9_]+", query.lower()) if len(t) > 1]
    if not terms:
        return []
    # Use OR semantics so natural-language questions still retrieve partial matches.
    match_query = " OR ".join(terms[:20])

    sql = (
        "SELECT c.id, c.path, c.rel_path, c.chunk_index, c.content, c.category, c.file_sha256, "
        "bm25(chunks_fts) as score "
        "FROM chunks_fts f "
        "JOIN chunks c ON c.id = f.rowid "
        "WHERE chunks_fts MATCH ? "
        "ORDER BY score LIMIT ?"
    )
    rows = conn.execute(sql, (match_query, top_k)).fetchall()
    result: List[Dict[str, object]] = []
    for r in rows:
        result.append(
            {
                "id": r[0],
                "path": r[1],
                "rel_path": r[2],
                "chunk_index": r[3],
                "content": r[4],
                "category": r[5],
                "file_sha256": r[6],
                "score": float(r[7]),
            }
        )
    return result


def build_prompt(question: str, chunks: Sequence[Dict[str, object]]) -> str:
    context_parts: List[str] = []
    for idx, c in enumerate(chunks, start=1):
        content = str(c["content"])[:1800]
        rel = str(c["rel_path"])
        chunk_index = c["chunk_index"]
        context_parts.append(f"[Source {idx}] {rel} (chunk {chunk_index})\n{content}")

    context = "\n\n---\n\n".join(context_parts)
    return (
        "You are a local retrieval assistant. Use only the provided context to answer. "
        "If the answer cannot be found in context, say that clearly and suggest what to search next.\n\n"
        f"Context:\n{context}\n\n"
        f"Question: {question}\n\n"
        "Answer with a concise response and cite source numbers like [Source 1]."
    )


def run_ollama(model: str, prompt: str) -> str:
    proc = subprocess.run(
        ["ollama", "run", model],
        input=prompt,
        text=True,
        capture_output=True,
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or "ollama run failed")
    return proc.stdout.strip()


def db_stats(conn: sqlite3.Connection) -> Dict[str, int]:
    files = conn.execute("SELECT COUNT(DISTINCT path) FROM chunks").fetchone()[0]
    chunks = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    return {"indexed_files": int(files), "indexed_chunks": int(chunks)}


def open_db(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db_path))
    ensure_schema(conn)
    return conn


def parse_common_args() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--config", default="rag_config.json", help="Path to JSON config file.")
    return parser


def resolve_paths(cfg: Dict[str, object]) -> Tuple[List[Path], List[Path], Path]:
    source_paths = [Path(p).expanduser() for p in cfg.get("source_paths", [])]
    exclude_paths = [Path(p).expanduser() for p in cfg.get("exclude_paths", [])]
    db_path = Path(str(cfg.get("database_path", "rag_index.sqlite3"))).expanduser()
    return source_paths, exclude_paths, db_path


def build_rel_path(path: Path, source_paths: Sequence[Path]) -> str:
    for src in source_paths:
        try:
            rel = path.resolve().relative_to(src.resolve())
            return str(rel)
        except Exception:
            continue
    return path.name


def safe_print(*args: object, **kwargs: object) -> None:
    print(*args, **kwargs)
